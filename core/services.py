"""Servicios de negocio de Bunny Finance.

Sin Qt: cualquier modulo (incluida la consola) puede usar esta capa. Cada
servicio expone metodos que ya no dependen de un widget, de modo que la vista
de Qt solo se ocupa de pintar el resultado.

Convenciones
------------
* Las rutas de base de datos se reciben por parametro y por defecto se
  resuelven con :mod:`core.paths` (absolutas, independientes del cwd).
* Los errores previstos se lanzan como :class:`ErrorServicio`; la capa superior
  decide como mostrarlos.
* El estado de una persona se guarda con los mismos codigos que usa la tabla
  ``personas``: 1 activado, 0 desactivado, 2 pendiente.
"""

from __future__ import annotations

import re
import sqlite3
from contextlib import closing
from dataclasses import dataclass, field
from datetime import datetime, time
from typing import Any, Iterable, Sequence

from core.paths import DATA_DB_PATH, FACES_DB_PATH

# --------------------------------------------------------------------------- #
# Constantes de dominio
# --------------------------------------------------------------------------- #

ESTATUS_TEXTO_A_CODIGO = {
    "activado": 1,
    "desactivado": 0,
    "pendiente": 2,
}
ESTATUS_CODIGO_A_TEXTO = {v: k for k, v in ESTATUS_TEXTO_A_CODIGO.items()}

FORMATO_HORARIO = "%H:%M"

DDL_PERSONAS = (
    "CREATE TABLE IF NOT EXISTS personas ("
    "id_persona INTEGER PRIMARY KEY AUTOINCREMENT, "
    "nombre_persona TEXT NOT NULL, "
    "imagen TEXT NOT NULL, "
    "encoding_perfil_izquierdo TEXT NOT NULL, "
    "encoding_frente TEXT NOT NULL, "
    "encoding_perfil_derecho TEXT NOT NULL, "
    "estatus INTEGER NOT NULL, "
    "created_at TEXT NOT NULL)"
)
DDL_SESIONES = (
    "CREATE TABLE IF NOT EXISTS sesiones ("
    "id_sesion INTEGER PRIMARY KEY AUTOINCREMENT, "
    "nombre TEXT NOT NULL, "
    "inicio TEXT NOT NULL, "
    "fin TEXT NOT NULL)"
)
DDL_ASISTENCIAS = (
    "CREATE TABLE IF NOT EXISTS asistencias ("
    "id_asistencia INTEGER PRIMARY KEY AUTOINCREMENT, "
    "id_persona INTEGER NOT NULL, "
    "id_sesion INTEGER NOT NULL, "
    "fecha TEXT NOT NULL, "
    "presente INTEGER NOT NULL)"
)
DDL_PERSONS = (
    "CREATE TABLE IF NOT EXISTS persons ("
    "id INTEGER PRIMARY KEY AUTOINCREMENT, "
    "name TEXT)"
)
DDL_ENCODINGS = (
    "CREATE TABLE IF NOT EXISTS encodings ("
    "id INTEGER PRIMARY KEY AUTOINCREMENT, "
    "person_id INTEGER, "
    "embedding BLOB)"
)


class ErrorServicio(Exception):
    """Error de negocio con mensaje apto para mostrar al usuario."""

    def __init__(self, mensaje: str, codigo: int = 1) -> None:
        super().__init__(mensaje)
        self.mensaje = mensaje
        self.codigo = codigo


# --------------------------------------------------------------------------- #
# Estructuras de datos
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class Persona:
    id: int
    nombre: str
    estatus: int
    created_at: str = ""

    @property
    def estatus_texto(self) -> str:
        return ESTATUS_CODIGO_A_TEXTO.get(self.estatus, f"desconocido({self.estatus})")


@dataclass(frozen=True)
class RosterEntry:
    """Fila de ``faces.db``: un rostro registrado con sus muestras."""

    id: int
    nombre: str
    muestras: int


@dataclass(frozen=True)
class Sesion:
    id: int
    nombre: str
    inicio: str
    fin: str

    @property
    def horario(self) -> str:
        return f"{self.inicio} - {self.fin}"


@dataclass(frozen=True)
class Asistencia:
    id: int
    id_persona: int
    id_sesion: int
    fecha: str
    presente: bool
    nombre: str = ""


@dataclass
class Diagnostico:
    personas_registradas: int = 0
    personas_activas: int = 0
    rostros: int = 0
    muestras: int = 0
    sesiones: int = 0
    asistencias: int = 0
    solo_en_rostros: list[str] = field(default_factory=list)
    solo_en_registro: list[str] = field(default_factory=list)
    nombres_duplicados: dict[str, int] = field(default_factory=dict)
    asistencias_duplicadas: list[tuple[int, int, int]] = field(default_factory=list)
    ruta_datos: str = ""
    ruta_rostros: str = ""


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #


def _conectar(ruta: str) -> sqlite3.Connection:
    try:
        return sqlite3.connect(ruta)
    except sqlite3.Error as exc:  # pragma: no cover - depende del entorno
        raise ErrorServicio(f"No se pudo abrir la base de datos {ruta}: {exc}") from exc


def normalizar_horario(valor: str, campo: str = "horario") -> str:
    """Valida y normaliza un horario a HH:MM (24 h)."""
    texto = str(valor or "").strip()
    if not re.fullmatch(r"\d{1,2}[:.]\d{2}", texto):
        raise ErrorServicio(f"{campo} invalido: {texto!r}. Usa el formato HH:MM (24 h).")
    horas, minutos = re.split(r"[:.]", texto)
    horas, minutos = int(horas), int(minutos)
    if not (0 <= horas <= 23 and 0 <= minutos <= 59):
        raise ErrorServicio(f"{campo} fuera de rango: {texto!r}")
    return f"{horas:02d}:{minutos:02d}"


def _hora_de_texto(valor: str) -> time | None:
    try:
        return datetime.strptime(valor, FORMATO_HORARIO).time()
    except (ValueError, TypeError):
        return None


# --------------------------------------------------------------------------- #
# Personas  (base de datos de datos)
# --------------------------------------------------------------------------- #


class PersonasService:
    """Gestion del registro de socios de la cooperativa (``personas``)."""

    def __init__(self, ruta: str = DATA_DB_PATH) -> None:
        self.ruta = ruta

    # -- esquema ---------------------------------------------------------- #
    def asegurar_esquema(self) -> None:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_PERSONAS)
            con.commit()

    # -- lectura ---------------------------------------------------------- #
    def listar(
        self,
        texto: str = "",
        estatus: str | None = None,
    ) -> list[Persona]:
        """Devuelve las personas del registro.

        A diferencia de ``DBManager.getPersonas``, no filtra por ``estatus = 1``:
        una persona desactivada tiene que seguir siendo visible para poder
        volver a activarla.
        """
        sql = (
            "SELECT id_persona, nombre_persona, estatus, created_at "
            "FROM personas WHERE 1=1"
        )
        params: list[Any] = []

        if texto:
            sql += " AND LOWER(nombre_persona) LIKE ?"
            params.append(f"%{texto.lower()}%")

        if estatus:
            codigo = ESTATUS_TEXTO_A_CODIGO.get(estatus.strip().lower())
            if codigo is None:
                raise ErrorServicio(
                    f"Estatus desconocido: {estatus!r}. "
                    f"Opciones: {', '.join(sorted(ESTATUS_TEXTO_A_CODIGO))}."
                )
            sql += " AND estatus = ?"
            params.append(codigo)

        sql += " ORDER BY nombre_persona COLLATE NOCASE"

        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(sql, params).fetchall()

        return [
            Persona(id=f[0], nombre=f[1], estatus=f[2], created_at=f[3] or "")
            for f in filas
        ]

    def obtener(self, id_persona: int) -> Persona | None:
        with closing(_conectar(self.ruta)) as con:
            fila = con.execute(
                "SELECT id_persona, nombre_persona, estatus, created_at "
                "FROM personas WHERE id_persona = ?",
                (int(id_persona),),
            ).fetchone()
        if not fila:
            return None
        return Persona(id=fila[0], nombre=fila[1], estatus=fila[2], created_at=fila[3] or "")

    def buscar_por_nombre(self, nombre: str) -> list[Persona]:
        return self.listar(texto=nombre)

    def por_nombre_exacto(self, nombre: str) -> Persona | None:
        objetivo = str(nombre or "").strip().casefold()
        coincidencias = [p for p in self.listar() if p.nombre.strip().casefold() == objetivo]
        return coincidencias[0] if coincidencias else None

    # -- escritura -------------------------------------------------------- #
    def establecer_estatus(self, id_persona: int, estatus: str) -> Persona:
        clave = str(estatus or "").strip().lower()
        codigo = ESTATUS_TEXTO_A_CODIGO.get(clave)
        if codigo is None:
            raise ErrorServicio(
                f"Estatus desconocido: {estatus!r}. "
                f"Opciones: {', '.join(sorted(ESTATUS_TEXTO_A_CODIGO))}."
            )

        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute(
                "UPDATE personas SET estatus = ? WHERE id_persona = ?",
                (codigo, int(id_persona)),
            )
            if cursor.rowcount == 0:
                existe = con.execute(
                    "SELECT 1 FROM personas WHERE id_persona = ?", (int(id_persona),)
                ).fetchone()
                if not existe:
                    raise ErrorServicio(f"No existe la persona con id {id_persona}.")
            con.commit()

        persona = self.obtener(id_persona)
        assert persona is not None
        return persona

    def eliminar(self, id_persona: int) -> None:
        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute("DELETE FROM personas WHERE id_persona = ?", (int(id_persona),))
            con.execute(
                "DELETE FROM asistencias WHERE id_persona = ?", (int(id_persona),)
            )
            con.commit()
        if cursor.rowcount == 0:
            raise ErrorServicio(f"No existe la persona con id {id_persona}.")

    def contar(self) -> dict[str, int]:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_PERSONAS)
            total = con.execute("SELECT COUNT(*) FROM personas").fetchone()[0]
            activos = con.execute(
                "SELECT COUNT(*) FROM personas WHERE estatus = 1"
            ).fetchone()[0]
        return {"total": total, "activos": activos}


# --------------------------------------------------------------------------- #
# Rostros  (faces.db, la base que usa el reconocimiento)
# --------------------------------------------------------------------------- #


class RostrosService:
    """Rostros y muestras de embedding registrados para el reconocimiento."""

    def __init__(self, ruta: str = FACES_DB_PATH) -> None:
        self.ruta = ruta

    def asegurar_esquema(self) -> None:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_PERSONS)
            con.execute(DDL_ENCODINGS)
            con.commit()

    def listar(self, texto: str = "") -> list[RosterEntry]:
        sql = (
            "SELECT p.id, p.name, COUNT(e.id) "
            "FROM persons p LEFT JOIN encodings e ON e.person_id = p.id "
            "WHERE 1=1"
        )
        params: list[Any] = []
        if texto:
            sql += " AND LOWER(p.name) LIKE ?"
            params.append(f"%{texto.lower()}%")
        sql += " GROUP BY p.id, p.name ORDER BY p.name COLLATE NOCASE, p.id"

        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(sql, params).fetchall()
        return [RosterEntry(id=f[0], nombre=f[1] or "(sin nombre)", muestras=f[2]) for f in filas]

    def obtener(self, id_person: int) -> RosterEntry | None:
        with closing(_conectar(self.ruta)) as con:
            fila = con.execute(
                "SELECT p.id, p.name, COUNT(e.id) FROM persons p "
                "LEFT JOIN encodings e ON e.person_id = p.id "
                "WHERE p.id = ? GROUP BY p.id, p.name",
                (int(id_person),),
            ).fetchone()
        if not fila:
            return None
        return RosterEntry(id=fila[0], nombre=fila[1] or "(sin nombre)", muestras=fila[2])

    def buscar_por_nombre(self, nombre: str) -> list[RosterEntry]:
        objetivo = str(nombre or "").strip().casefold()
        return [r for r in self.listar() if r.nombre.strip().casefold() == objetivo]

    def renombrar(self, id_person: int, nombre: str) -> RosterEntry:
        nuevo = str(nombre or "").strip()
        if not nuevo:
            raise ErrorServicio("El nombre no puede estar vacio.")
        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute(
                "UPDATE persons SET name = ? WHERE id = ?", (nuevo, int(id_person))
            )
            con.commit()
        if cursor.rowcount == 0:
            raise ErrorServicio(f"No existe el rostro con id {id_person}.")
        entrada = self.obtener(id_person)
        assert entrada is not None
        return entrada

    def eliminar(self, id_person: int, con_muestras: bool = True) -> None:
        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute("DELETE FROM persons WHERE id = ?", (int(id_person),))
            if con_muestras:
                con.execute("DELETE FROM encodings WHERE person_id = ?", (int(id_person),))
            con.commit()
        if cursor.rowcount == 0:
            raise ErrorServicio(f"No existe el rostro con id {id_person}.")

    def nombres_duplicados(self) -> dict[str, int]:
        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(
                "SELECT name, COUNT(*) FROM persons GROUP BY LOWER(name) HAVING COUNT(*) > 1"
            ).fetchall()
        return {(f[0] or "(sin nombre)"): f[1] for f in filas}

    def contar(self) -> dict[str, int]:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_PERSONS)
            con.execute(DDL_ENCODINGS)
            rostros = con.execute("SELECT COUNT(*) FROM persons").fetchone()[0]
            muestras = con.execute("SELECT COUNT(*) FROM encodings").fetchone()[0]
        return {"rostros": rostros, "muestras": muestras}


# --------------------------------------------------------------------------- #
# Sesiones
# --------------------------------------------------------------------------- #


class SesionesService:
    """CRUD de sesiones y deteccion de la sesion activa por reloj."""

    def __init__(self, ruta: str = DATA_DB_PATH) -> None:
        self.ruta = ruta

    def asegurar_esquema(self) -> None:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_SESIONES)
            con.commit()

    def listar(self) -> list[Sesion]:
        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(
                "SELECT id_sesion, nombre, inicio, fin FROM sesiones ORDER BY inicio, id_sesion"
            ).fetchall()
        return [Sesion(id=f[0], nombre=f[1], inicio=f[2], fin=f[3]) for f in filas]

    def obtener(self, id_sesion: int) -> Sesion | None:
        with closing(_conectar(self.ruta)) as con:
            fila = con.execute(
                "SELECT id_sesion, nombre, inicio, fin FROM sesiones WHERE id_sesion = ?",
                (int(id_sesion),),
            ).fetchone()
        return Sesion(id=fila[0], nombre=fila[1], inicio=fila[2], fin=fila[3]) if fila else None

    def obtener_por_nombre(self, nombre: str) -> Sesion | None:
        objetivo = str(nombre or "").strip().casefold()
        for sesion in self.listar():
            if sesion.nombre.strip().casefold() == objetivo:
                return sesion
        return None

    def crear(self, nombre: str, inicio: str, fin: str) -> Sesion:
        nombre = str(nombre or "").strip()
        if not nombre:
            raise ErrorServicio("El nombre de la sesion no puede estar vacio.")
        inicio = normalizar_horario(inicio, "Inicio")
        fin = normalizar_horario(fin, "Fin")

        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute(
                "INSERT INTO sesiones (nombre, inicio, fin) VALUES (?, ?, ?)",
                (nombre, inicio, fin),
            )
            con.commit()
            id_sesion = cursor.lastrowid
        return Sesion(id=id_sesion, nombre=nombre, inicio=inicio, fin=fin)

    def actualizar(
        self,
        id_sesion: int,
        nombre: str | None = None,
        inicio: str | None = None,
        fin: str | None = None,
    ) -> Sesion:
        actual = self.obtener(id_sesion)
        if actual is None:
            raise ErrorServicio(f"No existe la sesion con id {id_sesion}.")

        nombre = actual.nombre if nombre is None else str(nombre).strip()
        if not nombre:
            raise ErrorServicio("El nombre de la sesion no puede estar vacio.")
        inicio = actual.inicio if inicio is None else normalizar_horario(inicio, "Inicio")
        fin = actual.fin if fin is None else normalizar_horario(fin, "Fin")

        with closing(_conectar(self.ruta)) as con:
            con.execute(
                "UPDATE sesiones SET nombre = ?, inicio = ?, fin = ? WHERE id_sesion = ?",
                (nombre, inicio, fin, int(id_sesion)),
            )
            con.commit()
        return Sesion(id=int(id_sesion), nombre=nombre, inicio=inicio, fin=fin)

    def eliminar(self, id_sesion: int, con_asistencias: bool = True) -> None:
        with closing(_conectar(self.ruta)) as con:
            cursor = con.execute("DELETE FROM sesiones WHERE id_sesion = ?", (int(id_sesion),))
            if con_asistencias:
                con.execute(
                    "DELETE FROM asistencias WHERE id_sesion = ?", (int(id_sesion),)
                )
            con.commit()
        if cursor.rowcount == 0:
            raise ErrorServicio(f"No existe la sesion con id {id_sesion}.")

    def activa(self, momento: time | None = None) -> Sesion | None:
        """Sesion cuyo horario contiene a ``momento`` (inicio <= ahora < fin).

        Mismo criterio que el temporizador global de ``ui/main_window.py``.
        """
        ahora = momento or datetime.now().time()
        for sesion in self.listar():
            inicio = _hora_de_texto(sesion.inicio)
            fin = _hora_de_texto(sesion.fin)
            if inicio is None or fin is None:
                continue
            if inicio <= ahora < fin:
                return sesion
        return None


# --------------------------------------------------------------------------- #
# Asistencia
# --------------------------------------------------------------------------- #


class AsistenciaService:
    """Registro y consulta de asistencias por sesion."""

    def __init__(self, ruta: str = DATA_DB_PATH, personas: PersonasService | None = None) -> None:
        self.ruta = ruta
        self.personas = personas or PersonasService(ruta)

    def asegurar_esquema(self) -> None:
        with closing(_conectar(self.ruta)) as con:
            con.execute(DDL_ASISTENCIAS)
            con.commit()

    def registrar(
        self,
        id_persona: int,
        id_sesion: int,
        presente: bool = True,
        fecha: str | None = None,
    ) -> int:
        """Guarda (o actualiza) la asistencia de una persona en una sesion.

        Es idempotente: si ya habia una fila para (id_persona, id_sesion) la
        reemplaza en lugar de duplicarla, igual que hace la pantalla de
        asistencia al confirmar por segunda vez.
        """
        if self.personas.obtener(id_persona) is None:
            raise ErrorServicio(f"No existe la persona con id {id_persona}.")
        if SesionesService(self.ruta).obtener(id_sesion) is None:
            raise ErrorServicio(f"No existe la sesion con id {id_sesion}.")

        momento = fecha or datetime.now().isoformat()
        with closing(_conectar(self.ruta)) as con:
            con.execute(
                "DELETE FROM asistencias WHERE id_persona = ? AND id_sesion = ?",
                (int(id_persona), int(id_sesion)),
            )
            cursor = con.execute(
                "INSERT INTO asistencias (id_persona, id_sesion, fecha, presente) "
                "VALUES (?, ?, ?, ?)",
                (int(id_persona), int(id_sesion), momento, int(bool(presente))),
            )
            con.commit()
            return int(cursor.lastrowid)

    def registrar_nombre(
        self,
        nombre: str,
        id_sesion: int,
        presente: bool = True,
        fecha: str | None = None,
    ) -> int:
        """Registra la asistencia localizando a la persona por nombre.

        Devuelve el ``id_persona`` usado. Si el rostro existe en ``faces.db``
        pero la persona no esta en el registro de ``personas`` no se inventa un
        id: se avisa para que se registre primero.
        """
        persona = self.personas.por_nombre_exacto(nombre)
        if persona is not None:
            self.registrar(persona.id, id_sesion, presente, fecha)
            return persona.id

        raise ErrorServicio(
            f"No se encontro a {nombre!r} en el registro de personas. "
            "Registra primero el rostro desde la aplicacion."
        )

    def listar(
        self,
        id_sesion: int | None = None,
        id_persona: int | None = None,
    ) -> list[Asistencia]:
        sql = (
            "SELECT a.id_asistencia, a.id_persona, a.id_sesion, a.fecha, a.presente, "
            "       COALESCE(p.nombre_persona, '') "
            "FROM asistencias a LEFT JOIN personas p ON p.id_persona = a.id_persona "
            "WHERE 1=1"
        )
        params: list[Any] = []
        if id_sesion is not None:
            sql += " AND a.id_sesion = ?"
            params.append(int(id_sesion))
        if id_persona is not None:
            sql += " AND a.id_persona = ?"
            params.append(int(id_persona))
        sql += " ORDER BY a.id_sesion, p.nombre_persona COLLATE NOCASE, a.id_asistencia"

        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(sql, params).fetchall()

        return [
            Asistencia(
                id=f[0],
                id_persona=f[1],
                id_sesion=f[2],
                fecha=f[3],
                presente=bool(f[4]),
                nombre=f[5] or "",
            )
            for f in filas
        ]

    def eliminar(self, id_sesion: int, id_persona: int) -> None:
        with closing(_conectar(self.ruta)) as con:
            con.execute(
                "DELETE FROM asistencias WHERE id_sesion = ? AND id_persona = ?",
                (int(id_sesion), int(id_persona)),
            )
            con.commit()

    def asistencias_repetidas(self, id_sesion: int | None = None) -> list[tuple[int, int, int]]:
        """Grupos ``(id_persona, id_sesion, veces)`` con mas de una fila.

        La tabla ``asistencias`` no tiene restriccion de unicidad, asi que una
        misma persona puede quedar registrada varias veces en una sesion (por
        ejemplo si se pulsed "Exportar" o se reinicio el programa). Sirve para
        avisar antes de exportar.
        """
        sql = (
            "SELECT id_persona, id_sesion, COUNT(*) FROM asistencias "
            "GROUP BY id_persona, id_sesion HAVING COUNT(*) > 1"
        )
        params: list[Any] = []
        if id_sesion is not None:
            sql += " AND id_sesion = ?"
            params.append(int(id_sesion))
        with closing(_conectar(self.ruta)) as con:
            filas = con.execute(sql, params).fetchall()
        return [(f[0], f[1], f[2]) for f in filas]

    def limpiar_duplicados(self, id_sesion: int | None = None) -> int:
        """Deja una sola fila por persona y sesion (la mas reciente)."""
        repetidas = self.asistencias_repetidas(id_sesion)
        if not repetidas:
            return 0

        with closing(_conectar(self.ruta)) as con:
            for id_persona, id_sesion, _ in repetidas:
                con.execute(
                    "DELETE FROM asistencias WHERE id_persona = ? AND id_sesion = ? "
                    "AND id_asistencia NOT IN ("
                    "  SELECT MAX(id_asistencia) FROM asistencias "
                    "  WHERE id_persona = ? AND id_sesion = ?)",
                    (id_persona, id_sesion, id_persona, id_sesion),
                )
            con.commit()
        return sum(veces - 1 for _, _, veces in repetidas)

    def resumen(self, id_sesion: int) -> dict[str, int]:
        with closing(_conectar(self.ruta)) as con:
            presentes = con.execute(
                "SELECT COUNT(*) FROM asistencias WHERE id_sesion = ? AND presente = 1",
                (int(id_sesion),),
            ).fetchone()[0]
            ausentes = con.execute(
                "SELECT COUNT(*) FROM asistencias WHERE id_sesion = ? AND presente = 0",
                (int(id_sesion),),
            ).fetchone()[0]
        return {"presentes": presentes, "ausentes": ausentes, "total": presentes + ausentes}


# --------------------------------------------------------------------------- #
# Preferencias
# --------------------------------------------------------------------------- #

# clave de QSettings -> (etiqueta, tipo, valor por defecto, ayuda)
CLAVES_AJUSTES: dict[str, tuple[str, type, Any, str]] = {
    "theme/path": ("Tema", str, "styles/style.qss", "Hoja de estilos Qt"),
    "ui/language": ("Idioma", str, "Espanol", "Idioma de la interfaz"),
    "face/similarity_threshold": (
        "Umbral de similitud", float, 0.5, "Coseno minimo para aceptar un rostro (0-1)"
    ),
    "face/confirmations_needed": (
        "Confirmaciones", int, 2, "Detecciones consecutivas para confirmar"
    ),
    "face/detection_interval": (
        "Intervalo de deteccion", int, 5, "Frames entre inferencias"
    ),
    "camera/resolution": ("Resolucion", str, "640x480", "Ancho x Alto de la camara"),
    "camera/index": ("Indice de camara", int, 0, "0 = camara por defecto"),
    "attendance/required_seconds": (
        "Segundos requeridos", float, 2.0, "Presencia continua para marcar asistio"
    ),
    "attendance/ui_update_interval": (
        "Refresco de interfaz", float, 0.3, "Segundos entre actualizaciones de la tabla"
    ),
    "umap/n_neighbors": ("UMAP vecinos", int, 5, "n_neighbors de la proyeccion"),
    "umap/min_dist": ("UMAP min_dist", float, 0.1, "min_dist de la proyeccion"),
}


class AjustesService:
    """Lectura y escritura de las preferencias de la aplicacion.

    Usa el mismo QSettings que ``ui/app_settings.py`` para que la consola y la
    interfaz Compartan un unico almacen. QSettings es de QtCore y funciona sin
    QApplication, asi que se importa de forma perezosa.
    """

    def __init__(self, organizacion: str = "BunnyFinance", aplicacion: str = "BunnyFinance") -> None:
        self._organizacion = organizacion
        self._aplicacion = aplicacion
        self._store: Any | None = None

    @property
    def store(self) -> Any:
        if self._store is None:
            try:
                from PySide6.QtCore import QSettings
            except ImportError as exc:  # pragma: no cover - PySide6 es dependencia dura
                raise ErrorServicio(
                    "PySide6 no esta disponible, no se pueden leer las preferencias."
                ) from exc
            self._store = QSettings(self._organizacion, self._aplicacion)
        return self._store

    def claves(self) -> list[str]:
        return list(CLAVES_AJUSTES)

    def descripcion(self, clave: str) -> tuple[str, type, Any, str]:
        try:
            return CLAVES_AJUSTES[clave]
        except KeyError:
            raise ErrorServicio(
                f"Preferencia desconocida: {clave!r}. "
                f"Opciones: {', '.join(self.claves())}."
            ) from None

    def obtener(self, clave: str) -> Any:
        etiqueta, tipo, defecto, _ = self.descripcion(clave)
        del etiqueta
        valor = self.store.value(clave, defecto)
        try:
            if tipo is int:
                return int(float(valor))
            if tipo is float:
                return float(valor)
        except (TypeError, ValueError):
            return defecto
        return str(valor) if valor is not None else defecto

    def establecer(self, clave: str, valor: Any) -> Any:
        _, tipo, _, ayuda = self.descripcion(clave)
        convertido = self._validar(clave, tipo, valor, ayuda)
        self.store.setValue(clave, convertido)
        self.store.sync()
        return convertido

    def _validar(self, clave: str, tipo: type, valor: Any, ayuda: str) -> Any:
        texto = str(valor).strip()
        try:
            if tipo is int:
                convertido: Any = int(round(float(texto)))
            elif tipo is float:
                convertido = float(texto)
            else:
                convertido = texto
        except ValueError:
            raise ErrorServicio(
                f"{clave} espera un valor {tipo.__name__}, no {texto!r}. {ayuda}"
            ) from None

        if clave == "face/similarity_threshold" and not 0 < float(convertido) <= 1:
            raise ErrorServicio("El umbral de similitud debe estar entre 0 y 1.")
        if clave == "umap/min_dist" and not 0 <= float(convertido) <= 1:
            raise ErrorServicio("min_dist debe estar entre 0 y 1.")
        if clave == "camera/resolution" and not re.fullmatch(r"\d+x\d+", texto):
            raise ErrorServicio("La resolucion debe tener el formato ANCHOxALTO (640x480).")

        return convertido

    def listar(self) -> dict[str, Any]:
        return {clave: self.obtener(clave) for clave in self.claves()}

    def vista(self) -> list[tuple[str, str, Any, str]]:
        """Filas ``(clave, etiqueta, valor, ayuda)`` para pintar en consola."""
        filas = []
        for clave, (etiqueta, _, defecto, ayuda) in CLAVES_AJUSTES.items():
            filas.append((clave, etiqueta, self.obtener(clave) or defecto, ayuda))
        return filas

    def restaurar(self) -> None:
        for clave, (_, tipo, defecto, _) in CLAVES_AJUSTES.items():
            self.store.setValue(clave, defecto)
        self.store.sync()


# --------------------------------------------------------------------------- #
# Exportacion a Excel
# --------------------------------------------------------------------------- #

CABECERA_ASISTENCIA = ["ID", "Nombre", "Similitud", "Asistencia"]


class ExcelService:
    """Escritura del .xlsx de asistencia.

    Acepta filas ya construidas para que tanto la vista de Qt (que las saca de
    su tabla) como la consola (que las saca de la base de datos) usen el mismo
    generador de archivo.
    """

    @staticmethod
    def exportar(
        filas: Iterable[Sequence[Any]],
        ruta: str,
        titulo: str = "Asistencias",
    ) -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
        except ImportError as exc:  # pragma: no cover
            raise ErrorServicio(
                "Falta openpyxl; instalalo con: pip install openpyxl"
            ) from exc

        libro = Workbook()
        hoja = libro.active
        hoja.title = titulo[:31] or "Asistencias"

        for indice, titulo_columna in enumerate(CABECERA_ASISTENCIA, start=1):
            hoja.cell(row=1, column=indice, value=titulo_columna)

        for numero, fila in enumerate(filas, start=2):
            for indice, valor in enumerate(fila, start=1):
                hoja.cell(row=numero, column=indice, value="" if valor is None else valor)

        for indice in range(1, len(CABECERA_ASISTENCIA) + 1):
            hoja.column_dimensions[get_column_letter(indice)].auto_size = True

        libro.save(ruta)
        return ruta


# --------------------------------------------------------------------------- #
# Fachada
# --------------------------------------------------------------------------- #


class Core:
    """Punto de entrada unico a la logica de negocio.

        core = Core()
        core.sesiones.crear("Algebra III", "08:00", "09:00")
    """

    def __init__(self, ruta_datos: str = DATA_DB_PATH, ruta_rostros: str = FACES_DB_PATH) -> None:
        self.ruta_datos = ruta_datos
        self.ruta_rostros = ruta_rostros

        self.personas = PersonasService(ruta_datos)
        self.sesiones = SesionesService(ruta_datos)
        self.asistencias = AsistenciaService(ruta_datos, self.personas)
        self.rostros = RostrosService(ruta_rostros)
        self.ajustes = AjustesService()
        self.excel = ExcelService()

    def asegurar_esquema(self) -> None:
        """Crea las tablas que falten. Idempotente."""
        self.personas.asegurar_esquema()
        self.sesiones.asegurar_esquema()
        self.asistencias.asegurar_esquema()
        self.rostros.asegurar_esquema()

    # -- utilidades compuestas ------------------------------------------- #
    def resolver_sesion(self, referencia: str | int | None) -> Sesion:
        """Resuelve un id, un nombre o ``"activa"`` a una sesion concreta."""
        if referencia is None or referencia == "" or referencia == "activa":
            sesion = self.sesiones.activa()
            if sesion is None:
                raise ErrorServicio(
                    "No hay ninguna sesion activa en este momento. "
                    "Indica un id o un nombre de sesion."
                )
            return sesion

        if isinstance(referencia, int):
            sesion = self.sesiones.obtener(referencia)
        elif str(referencia).strip().isdigit():
            sesion = self.sesiones.obtener(int(referencia))
        else:
            sesion = self.sesiones.obtener_por_nombre(referencia)

        if sesion is None:
            raise ErrorServicio(f"No se encontro la sesion {referencia!r}.")
        return sesion

    def filas_asistencia(self, id_sesion: int) -> list[tuple[Any, ...]]:
        """Filas listas para :meth:`ExcelService.exportar`.

        Incluye a todo el roster conocido, no solo a quien asistio, para que los
        ausentes figuren explicitamente como "No asistio" (igual que la tabla
        de la pantalla de asistencia).
        """
        registros = {a.id_persona: a for a in self.asistencias.listar(id_sesion=id_sesion)}

        def estado(id_persona: int) -> str:
            registro = registros.get(id_persona)
            return "Asistio" if registro and registro.presente else "No asistio"

        filas: list[tuple[Any, ...]] = []
        nombres_registro: set[str] = set()

        for persona in self.personas.listar():
            nombres_registro.add(persona.nombre.strip().casefold())
            filas.append((persona.id, persona.nombre, "", estado(persona.id)))

        # Rostros dados de alta en faces.db que aun no estan en el registro: se
        # listan para que el Excel cuadre con lo que ve el operador.
        for entrada in self.rostros.listar():
            if entrada.nombre.strip().casefold() in nombres_registro:
                continue
            filas.append((entrada.id, entrada.nombre, "", estado(entrada.id)))

        return filas

    def exportar_sesion(self, id_sesion: int, ruta: str | None = None) -> str:
        sesion = self.sesiones.obtener(id_sesion)
        if sesion is None:
            raise ErrorServicio(f"No existe la sesion con id {id_sesion}.")

        if ruta is None:
            seguro = sesion.nombre.replace(" ", "_")
            ruta = f"{seguro}_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

        return self.excel.exportar(self.filas_asistencia(id_sesion), ruta, titulo="Asistencias")

    def diagnostico(self) -> Diagnostico:
        """Resumen del estado de las dos bases de datos y sus discrepancias."""
        conteo_personas = self.personas.contar()
        conteo_rostros = self.rostros.contar()

        with closing(_conectar(self.ruta_datos)) as con:
            con.execute(DDL_SESIONES)
            con.execute(DDL_ASISTENCIAS)
            sesiones = con.execute("SELECT COUNT(*) FROM sesiones").fetchone()[0]
            asistencias = con.execute("SELECT COUNT(*) FROM asistencias").fetchone()[0]

        nombres_registro = {p.nombre.strip().casefold() for p in self.personas.listar()}
        nombres_rostros = {r.nombre.strip().casefold() for r in self.rostros.listar()}

        return Diagnostico(
            personas_registradas=conteo_personas["total"],
            personas_activas=conteo_personas["activos"],
            rostros=conteo_rostros["rostros"],
            muestras=conteo_rostros["muestras"],
            sesiones=sesiones,
            asistencias=asistencias,
            solo_en_rostros=sorted(nombres_rostros - nombres_registro),
            solo_en_registro=sorted(nombres_registro - nombres_rostros),
            nombres_duplicados=self.rostros.nombres_duplicados(),
            asistencias_duplicadas=self.asistencias.asistencias_repetidas(),
            ruta_datos=self.ruta_datos,
            ruta_rostros=self.ruta_rostros,
        )
