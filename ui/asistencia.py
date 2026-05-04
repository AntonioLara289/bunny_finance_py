from PySide6.QtWidgets import (
    QLineEdit, QTableWidget, QTableWidgetItem, QComboBox,
    QLabel, QVBoxLayout, QFrame, QHBoxLayout, QHeaderView,
    QSizePolicy, QFileDialog, QPushButton
)
from PySide6 import QtWidgets, QtCore
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from ui.camaraInsightFaceWidget import CameraInsightFaceWidget, DB_PATH
from database.db_manager import DBManager
import sqlite3
import time
from datetime import datetime
from log import log

asistencias_status = {1: "Asistió", 0: "No asistió"}

class AsistenciaPantalla(QtWidgets.QWidget):
    def __init__(self, session_name=None, auto_start_camera=False, session_id=None):
        super().__init__()

        self.setAccessibleName("AsistenciaPantalla")

        # Configuracion
        self.SIMILARITY_THRESHOLD = 60.0
        self.REQUIRED_SECONDS = 5.0

        # Variables de rendimiento
        self.UI_UPDATE_INTERVAL = 0.3  # segundos
        self.TIMER_GRACE_PERIOD = 1.0  # segundos

        # person_id -> start_time
        self.recognition_timers = {}
        self.attendance_given = set()
        self.last_ui_update = {}

        # Interfaz
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(20, 20, 20, 20)
        main_layout.setSpacing(15)

        title = QLabel("Registro de Asistencia")
        title.setAlignment(QtCore.Qt.AlignCenter)
        main_layout.addWidget(title)

        # Sesión info
        session_info = QLabel(f"Sesión: {session_name if session_name else 'Sin sesión'}")
        session_info.setAlignment(QtCore.Qt.AlignCenter)
        session_info_font = session_info.font()
        session_info_font.setPointSize(10)
        session_info_font.setBold(True)
        session_info.setFont(session_info_font)
        main_layout.addWidget(session_info)

        # Camara
        cam_frame = QFrame()
        cam_layout = QVBoxLayout(cam_frame)


        self.dbManager = DBManager()
        self.session_name = session_name
        self.session_id = session_id

        self.face_conn = sqlite3.connect(DB_PATH)
        self.face_cursor = self.face_conn.cursor()
        self.face_cursor.execute("""
            SELECT persons.id, persons.name
            FROM persons
        """)
        self.getPersonas = self.face_cursor.fetchall()

        self.camera_widget = CameraInsightFaceWidget()
        self.camera_widget.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.camera_widget.setMinimumHeight(320)
        cam_layout.addWidget(self.camera_widget)
        self.camera_widget.personaConfirmada.connect(self.actualizarAsistencia)
        main_layout.addWidget(cam_frame)
        if auto_start_camera:
            try:
                self.camera_widget.start_camera()
            except Exception as e:
                print(f"[ERROR] No se pudo iniciar la cámara automáticamente: {e}")

        # Tabla
        table_frame = QFrame()
        table_layout = QVBoxLayout(table_frame)

        table_label = QLabel("Lista de personas")
        table_layout.addWidget(table_label)

        num_rows = len(self.getPersonas)
        self.table = QTableWidget(num_rows, 4)
        self.table.setHorizontalHeaderLabels([
            "ID", "Nombre", "Similitud", "Asistencia"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)

        # ID → ROW MAP (Busqueda O(1))
        self.id_to_row = {}

        for row_idx, (persona_id, persona_name) in enumerate(self.getPersonas):
            try:
                self.id_to_row[persona_id] = row_idx

                self.table.setItem(row_idx, 0, QTableWidgetItem(str(persona_id)))
                self.table.setItem(row_idx, 1, QTableWidgetItem(str(persona_name)))
                self.table.setItem(row_idx, 2, QTableWidgetItem("0%"))

                combo = QComboBox()
                combo.addItems(["No asistió", "Asistió"])
                self.table.setCellWidget(row_idx, 3, combo)
            except Exception as e:
                print(f"[ERROR] Error llenando la tabla de asistencia en fila {row_idx}: {e}")

        table_layout.addWidget(self.table)
        main_layout.addWidget(table_frame)

        # Exportacion
        export_btn = QPushButton("Exportar a Excel")
        export_btn.setMinimumHeight(40)
        export_btn.clicked.connect(
            lambda: (self.exportarExcel(), log.push("Lista exportada"))
        )
        try:
            export_btn.setToolTip("Exportar la lista actual de asistencias a Excel")
            export_btn.setAccessibleName("ExportarExcel")
            export_btn.setIcon(self.style().standardIcon(QtWidgets.QStyle.SP_DialogSaveButton))
        except Exception:
            pass
        main_layout.addWidget(export_btn, alignment=QtCore.Qt.AlignRight)

    # Logica de asistencia
    def actualizarAsistencia(self, name, id_, similarity):
        if id_ in self.attendance_given:
            return

        now = time.time()

        last_ui = self.last_ui_update.get(id_, 0)
        if now - last_ui < self.UI_UPDATE_INTERVAL:
            return
        self.last_ui_update[id_] = now

        row = self.id_to_row.get(id_)
        if row is None:
            return

        similarity_item = self.table.item(row, 2)
        similarity_item.setText(f"{similarity:.1f}%")

        if similarity >= self.SIMILARITY_THRESHOLD:
            if id_ not in self.recognition_timers:
                self.recognition_timers[id_] = now
            else:
                elapsed = now - self.recognition_timers[id_]
                if elapsed >= self.REQUIRED_SECONDS:
                    combo = self.table.cellWidget(row, 3)
                    if combo:
                        combo.setCurrentText("Asistió")
                        self.attendance_given.add(id_)
                        log.push(f"Asistencia registrada para {name}")
        else:
            start = self.recognition_timers.get(id_)
            if start and now - start > self.TIMER_GRACE_PERIOD:
                del self.recognition_timers[id_]

    # Exportacion
    def exportarExcel(self):
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Guardar archivo Excel",
            "asistencia.xlsx",
            "Excel Files (*.xlsx)"
        )
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencias"

        for col in range(self.table.columnCount()):
            header = self.table.horizontalHeaderItem(col).text()
            ws.cell(row=1, column=col + 1, value=header)

        for row in range(self.table.rowCount()):
            for col in range(self.table.columnCount()):
                item = self.table.item(row, col)
                if item:
                    ws.cell(row=row + 2, column=col + 1, value=item.text())
                else:
                    widget = self.table.cellWidget(row, col)
                    if isinstance(widget, QComboBox):
                        ws.cell(
                            row=row + 2,
                            column=col + 1,
                            value=widget.currentText()
                        )

        for col in range(1, self.table.columnCount() + 1):
            ws.column_dimensions[get_column_letter(col)].auto_size = True

        wb.save(path)

    def exportarExcelTo(self, path):
        # Same logic as exportarExcel but writing directly to `path`
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Asistencias"

        for col in range(self.table.columnCount()):
            header = self.table.horizontalHeaderItem(col).text()
            ws.cell(row=1, column=col + 1, value=header)

        for row in range(self.table.rowCount()):
            for col in range(self.table.columnCount()):
                item = self.table.item(row, col)
                if item:
                    ws.cell(row=row + 2, column=col + 1, value=item.text())
                else:
                    widget = self.table.cellWidget(row, col)
                    if isinstance(widget, QComboBox):
                        ws.cell(
                            row=row + 2,
                            column=col + 1,
                            value=widget.currentText()
                        )

        for col in range(1, self.table.columnCount() + 1):
            ws.column_dimensions[get_column_letter(col)].auto_size = True

        wb.save(path)

    def closeEvent(self, event):
        try:
            if getattr(self, 'camera_widget', None):
                try:
                    self.camera_widget.stop_camera()
                except Exception:
                    pass

            if getattr(self, 'face_conn', None):
                try:
                    self.face_conn.close()
                except Exception:
                    pass

            if self.session_name:
                fecha = datetime.now().strftime('%Y-%m-%d')
                safe_name = self.session_name.replace(' ', '_')
                filename = f"{safe_name}_{fecha}.xlsx"
                try:
                    self.exportarExcelTo(filename)
                except Exception:
                    pass
        finally:
            event.accept()
